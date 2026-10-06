import hashlib
import json
import re
from pathlib import Path
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "build"
OUT.mkdir(exist_ok=True)
slides = json.loads((ROOT / "presentation/slides.json").read_text())
report = (ROOT / "docs/PROJECT_REPORT.md").read_text()

def textbox(slide, x, y, w, h, text, size, color, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(18)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
for i, content in enumerate(slides, 1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("0B1220")
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(.58), Inches(.08), Inches(.45))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string("2DD4BF")
    bar.line.fill.background()
    textbox(slide, .95, .58, 11.5, .4, content["tag"].upper(), 13, "2DD4BF", True)
    textbox(slide, .95, 1.25, 11.5, .9, content["title"], 34, "F8FAFC", True)
    textbox(slide, 1.0, 2.6, 11.2, 3.8, content["body"], 24, "CBD5E1")
    textbox(slide, .95, 6.95, 11.0, .3, "DATAPILOT AI | PROTOTYPE | 07 OCT 2026", 10, "94A3B8")
    textbox(slide, 12.0, 6.95, .6, .3, str(i).zfill(2), 10, "2DD4BF")
    slide.notes_slide.notes_text_frame.text = content["notes"]
prs.save(OUT / "DataPilot_AI_Presentation.pptx")

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="ReportBody", fontName="Helvetica", fontSize=10, leading=15, spaceAfter=12))
styles.add(ParagraphStyle(name="ReportHeading", fontName="Helvetica-Bold", fontSize=23, leading=29, textColor=colors.HexColor("#0B1220"), spaceAfter=18))
styles.add(ParagraphStyle(name="ReportTag", fontName="Helvetica-Bold", fontSize=10, leading=14, textColor=colors.HexColor("#087F8C"), spaceAfter=14))
blocks = re.split(r"^## ", report, flags=re.MULTILINE)
intro = blocks[0].splitlines()
story = [Spacer(1, 70), Paragraph("ENGINEERING CASE STUDY", styles["ReportTag"]), Paragraph("DataPilot AI", styles["ReportHeading"]), Paragraph("Conversational spreadsheet analysis and Gmail automation", styles["ReportBody"]), Spacer(1, 30), Paragraph("Owner: AnyalstRushi | Prepared 7 October 2026", styles["ReportBody"]), Paragraph("Evidence-backed prototype. Original workflow export, dynamic report receipt, and production controls remain pending.", styles["ReportBody"])]
for block in blocks[1:]:
    title, _, body = block.partition("\n")
    story += [PageBreak(), Paragraph("PROJECT REPORT | PROTOTYPE", styles["ReportTag"]), Paragraph(escape(title), styles["ReportHeading"])]
    for paragraph in body.strip().split("\n\n"):
        story.append(Paragraph(escape(paragraph).replace("\n", "<br/>"), styles["ReportBody"]))

def footer(canvas, doc):
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#526174"))
    canvas.drawString(48, 30, "DATAPILOT AI | EVIDENCE-BACKED PROTOTYPE")
    canvas.drawRightString(A4[0] - 48, 30, str(doc.page))

SimpleDocTemplate(str(OUT / "DataPilot_AI_Project_Report.pdf"), pagesize=A4, leftMargin=48, rightMargin=48, topMargin=45, bottomMargin=55, title="DataPilot AI Project Report", author="AnyalstRushi").build(story, onFirstPage=footer, onLaterPages=footer)
doc = Document()
doc.add_heading("DataPilot AI", 0)
doc.add_paragraph("Engineering project report | Evidence-backed prototype")
doc.add_paragraph("Owner: AnyalstRushi | Prepared 7 October 2026")
for block in blocks[1:]:
    title, _, body = block.partition("\n")
    doc.add_page_break()
    doc.add_heading(title, 1)
    for paragraph in body.strip().split("\n\n"):
        doc.add_paragraph(paragraph)
doc.save(OUT / "DataPilot_AI_Project_Report.docx")
manifest = {p.name: {"bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in OUT.iterdir() if p.is_file() and p.name != "MANIFEST.json"}
if len(manifest) != 3 or not all(v["bytes"] > 0 for v in manifest.values()):
    raise RuntimeError("Missing or empty project deliverable")
(OUT / "MANIFEST.json").write_text(json.dumps(manifest, indent=2))
print(json.dumps({"slides": len(slides), "deliverables": manifest}, indent=2))
